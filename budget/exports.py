"""Exports Excel (openpyxl) et PDF (reportlab)."""

from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from . import services

HEADER_FILL = PatternFill('solid', fgColor='0F766E')
HEADER_FONT = Font(color='FFFFFF', bold=True)
TEAL = colors.HexColor('#0f766e')


def _currency(user):
    profile = getattr(user, 'profile', None)
    return profile.currency if profile else 'FCFA'


def _stamp():
    return timezone.localtime().strftime('%Y-%m-%d_%H%M')


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def expenses_xlsx(user, queryset):
    devise = _currency(user)
    wb = Workbook()

    ws = wb.active
    ws.title = 'Depenses'
    headers = ['Date', 'Categorie', 'Montant', 'Mode de paiement', 'Description', 'Note']
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal='center')

    total = Decimal('0')
    for expense in queryset:
        total += expense.amount
        ws.append([
            expense.date,
            expense.category.name,
            float(expense.amount),
            expense.get_payment_method_display(),
            expense.description,
            expense.note,
        ])

    ws.append([])
    row = ws.max_row + 1
    ws.cell(row=row, column=2, value='TOTAL').font = Font(bold=True)
    ws.cell(row=row, column=3, value=float(total)).font = Font(bold=True)

    ws.freeze_panes = 'A2'
    for i, width in enumerate([12, 18, 14, 20, 40, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).number_format = f'#,##0.00 "{devise}"'

    # Feuille de synthese par categorie sur le mois en cours
    ws2 = wb.create_sheet('Synthese du mois')
    ws2.append(['Categorie', 'Total', 'Part (%)', 'Nombre'])
    for col in range(1, 5):
        cell = ws2.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    month = services.period_bounds('month')
    for row_data in services.breakdown_by_category(user, month):
        ws2.append([
            row_data['name'], float(row_data['total']),
            round(row_data['share'], 2), row_data['count'],
        ])
    for i, width in enumerate([24, 16, 12, 10], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = width

    # Feuille des revenus : sans elle, l'export ne dirait rien du solde.
    ws3 = wb.create_sheet('Revenus')
    ws3.append(['Date', 'Source', 'Montant', 'Mode de reception', 'Recurrent',
                'Description'])
    for col in range(1, 7):
        cell = ws3.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    encaisse = Decimal('0')
    for income in user.incomes.select_related('source').all():
        encaisse += income.amount
        ws3.append([
            income.date, income.source.name, float(income.amount),
            income.get_method_display(), 'oui' if income.is_recurring else 'non',
            income.description,
        ])
    ws3.append([])
    ligne = ws3.max_row + 1
    ws3.cell(row=ligne, column=2, value='TOTAL ENCAISSE').font = Font(bold=True)
    ws3.cell(row=ligne, column=3, value=float(encaisse)).font = Font(bold=True)
    # Les depenses de la feuille 1 sont filtrees, pas les revenus : on le dit,
    # plutot que d'afficher un solde dont le perimetre serait ambigu.
    ws3.cell(
        row=ligne + 1, column=2,
        value='SOLDE (total encaisse - depenses exportees)',
    ).font = Font(bold=True)
    ws3.cell(row=ligne + 1, column=3, value=float(encaisse - total)).font = Font(bold=True)
    ws3.freeze_panes = 'A2'
    for i, width in enumerate([12, 20, 14, 20, 12, 40], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="budget-control_{_stamp()}.xlsx"'
    )
    return response


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def expenses_pdf(user, queryset, kind='month'):
    devise = _currency(user)
    period = services.period_bounds(kind)
    comparison = services.compare(user, kind)
    breakdown = services.breakdown_by_category(user, period)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=1.6 * cm, rightMargin=1.6 * cm,
        topMargin=1.6 * cm, bottomMargin=1.6 * cm,
        title='Budget Control - Rapport',
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle('t', parent=styles['Title'], textColor=TEAL, fontSize=20)
    muted = ParagraphStyle('m', parent=styles['Normal'],
                           textColor=colors.HexColor('#64748b'), fontSize=9)

    story = [
        Paragraph('Budget Control', title),
        Paragraph(
            f'Rapport {services.PERIOD_LABELS[kind][0].lower()} — '
            f'{period.start:%d/%m/%Y} au {period.end:%d/%m/%Y}', muted,
        ),
        Paragraph(f'Utilisateur : {user.username} — genere le '
                  f'{timezone.localtime():%d/%m/%Y a %H:%M}', muted),
        Spacer(1, 0.8 * cm),
    ]

    total = sum((e.amount for e in queryset), Decimal('0'))
    synthese = [
        ['Total periode', f'{comparison["current"]:,.2f} {devise}'],
        ['Periode precedente', f'{comparison["previous"]:,.2f} {devise}'],
        ['Variation',
         f'{comparison["variation"]:+.1f} %' if comparison['variation'] is not None
         else 'pas de reference'],
        ['Categorie principale', breakdown[0]['name'] if breakdown else '—'],
        ['Depenses exportees', f'{queryset.count()} ({total:,.2f} {devise})'],
    ]
    table = Table(synthese, colWidths=[6 * cm, 10.4 * cm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f1f5f9')),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#475569')),
        ('FONTNAME', (1, 0), (1, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#e2e8f0')),
        ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    story += [table, Spacer(1, 0.8 * cm),
              Paragraph('Repartition par categorie', styles['Heading3'])]

    if breakdown:
        rows = [['Categorie', 'Montant', 'Part']]
        rows += [
            [r['name'], f'{r["total"]:,.2f} {devise}', f'{r["share"]:.1f} %']
            for r in breakdown
        ]
        cat_table = Table(rows, colWidths=[8 * cm, 5 * cm, 3.4 * cm])
        cat_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), TEAL),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#e2e8f0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1),
             [colors.white, colors.HexColor('#f8fafc')]),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(cat_table)
    else:
        story.append(Paragraph('Aucune depense sur cette periode.', muted))

    story += [Spacer(1, 0.8 * cm), Paragraph('Detail des depenses', styles['Heading3'])]
    detail = [['Date', 'Categorie', 'Montant', 'Paiement', 'Description']]
    for expense in queryset[:400]:
        detail.append([
            expense.date.strftime('%d/%m/%y'),
            expense.category.name,
            f'{expense.amount:,.2f}',
            expense.get_payment_method_display(),
            Paragraph(expense.description or '—', muted),
        ])
    detail_table = Table(detail, colWidths=[2.2 * cm, 3.6 * cm, 3 * cm, 3 * cm, 4.6 * cm],
                         repeatRows=1)
    detail_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (2, 1), (2, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(detail_table)
    if queryset.count() > 400:
        story.append(Paragraph(
            f'… {queryset.count() - 400} depenses supplementaires non listees. '
            'Utilisez l’export Excel pour le detail complet.', muted))

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="budget-control_{kind}_{_stamp()}.pdf"'
    )
    return response
