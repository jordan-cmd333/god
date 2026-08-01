"""Tests des regles de gestion et des vues principales."""

from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from . import services
from .models import (
    Alert, BudgetLimit, Category, Expense, Income, IncomeSource, Report,
)


class BaseCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('alice', password='motdepasse-123')
        self.food = Category.objects.get(user=self.user, name='Nourriture')
        self.transport = Category.objects.get(user=self.user, name='Transport')
        self.today = date(2026, 7, 15)  # un mercredi

    def spend(self, amount, when=None, category=None):
        return Expense.objects.create(
            user=self.user,
            category=category or self.food,
            amount=Decimal(amount),
            date=when or self.today,
        )


class BootstrapTests(BaseCase):
    def test_les_categories_par_defaut_sont_creees(self):
        self.assertEqual(Category.objects.filter(user=self.user).count(), 11)

    def test_le_profil_est_cree(self):
        self.assertEqual(self.user.profile.alert_threshold, 80)


class PeriodTests(BaseCase):
    def test_bornes_de_semaine_du_lundi_au_dimanche(self):
        week = services.period_bounds('week', self.today)
        self.assertEqual(week.start, date(2026, 7, 13))
        self.assertEqual(week.end, date(2026, 7, 19))

    def test_bornes_de_mois(self):
        month = services.period_bounds('month', self.today)
        self.assertEqual((month.start, month.end), (date(2026, 7, 1), date(2026, 7, 31)))

    def test_periode_precedente(self):
        self.assertEqual(services.previous_period('day', self.today).start,
                         date(2026, 7, 14))
        self.assertEqual(services.previous_period('month', self.today).start,
                         date(2026, 6, 1))


class AggregationTests(BaseCase):
    def test_les_depenses_sont_additionnees_par_periode(self):
        self.spend('1000')
        self.spend('500')
        self.spend('9999', when=self.today - timedelta(days=40))  # hors mois

        day = services.period_bounds('day', self.today)
        month = services.period_bounds('month', self.today)
        self.assertEqual(services.total_for(self.user, day), Decimal('1500'))
        self.assertEqual(services.total_for(self.user, month), Decimal('1500'))

    def test_categorie_la_plus_depensiere(self):
        self.spend('1000', category=self.food)
        self.spend('4000', category=self.transport)
        top = services.top_category(self.user, services.period_bounds('month', self.today))
        self.assertEqual(top['name'], 'Transport')
        self.assertAlmostEqual(top['share'], 80.0)

    def test_pas_de_pourcentage_sans_periode_de_reference(self):
        self.spend('1500')
        result = services.compare(self.user, 'day', self.today)
        self.assertIsNone(result['variation'])
        self.assertEqual(result['previous'], Decimal('0.00'))

    def test_comparaison_entre_periodes(self):
        self.spend('1000', when=self.today - timedelta(days=1))
        self.spend('1500', when=self.today)
        result = services.compare(self.user, 'day', self.today)
        self.assertEqual(result['current'], Decimal('1500'))
        self.assertEqual(result['previous'], Decimal('1000'))
        self.assertAlmostEqual(result['variation'], 50.0)
        self.assertFalse(result['improving'])


class AlertTests(BaseCase):
    def setUp(self):
        super().setUp()
        self.limit = BudgetLimit.objects.create(
            user=self.user, period='month', amount=Decimal('10000')
        )

    def test_aucune_alerte_sous_le_seuil(self):
        self.spend('7000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(Alert.objects.count(), 0)

    def test_alerte_a_80_pourcent(self):
        self.spend('8000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(Alert.objects.filter(level=Alert.Level.WARNING).count(), 1)
        self.assertEqual(Alert.objects.filter(level=Alert.Level.EXCEEDED).count(), 0)

    def test_alerte_de_depassement(self):
        self.spend('11000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(Alert.objects.filter(level=Alert.Level.EXCEEDED).count(), 1)

    def test_les_alertes_ne_sont_pas_dupliquees(self):
        self.spend('8500')
        services.evaluate_alerts(self.user, self.today)
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(Alert.objects.count(), 1)

    def test_une_seule_alerte_affichee_par_limite(self):
        # Un depassement enregistre les deux niveaux, mais n'en affiche qu'un.
        self.spend('12000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(Alert.objects.count(), 2)
        affichees = services.unread_alerts(self.user)
        self.assertEqual(len(affichees), 1)
        self.assertEqual(affichees[0].level, Alert.Level.EXCEEDED)

    def test_le_seuil_personnalise_est_respecte(self):
        self.user.profile.alert_threshold = 50
        self.user.profile.save()
        self.spend('5500')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(Alert.objects.filter(level=Alert.Level.WARNING).count(), 1)

    def test_limite_par_categorie_ignore_les_autres_categories(self):
        limit = BudgetLimit.objects.create(
            user=self.user, period='month', category=self.food, amount=Decimal('1000')
        )
        self.spend('5000', category=self.transport)
        status = services.limit_status(self.user, limit, self.today)
        self.assertEqual(status['spent'], Decimal('0'))
        self.assertIsNone(status['level'])

    def test_barre_de_progression_plafonnee_a_100(self):
        self.spend('30000')
        status = services.limit_status(self.user, self.limit, self.today)
        self.assertEqual(status['bar_percent'], 100)
        self.assertEqual(status['percent'], 300)
        self.assertEqual(status['remaining'], Decimal('-20000'))


class IncomeTests(BaseCase):
    """Les revenus suivent le meme concept que les depenses."""

    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)
        self.salaire = IncomeSource.objects.get(user=self.user, name='Salaire')
        self.vente = IncomeSource.objects.get(user=self.user, name='Vente')

    def earn(self, amount, when=None, source=None):
        return Income.objects.create(
            user=self.user, source=source or self.salaire,
            amount=Decimal(amount), date=when or self.today,
        )

    def test_les_sources_par_defaut_sont_creees(self):
        self.assertEqual(IncomeSource.objects.filter(user=self.user).count(), 10)

    def test_les_revenus_sont_additionnes_par_periode(self):
        self.earn('180000')
        self.earn('20000')
        self.earn('99999', when=self.today - timedelta(days=40))
        month = services.period_bounds('month', self.today)
        self.assertEqual(services.income_total(self.user, month), Decimal('200000'))

    def test_source_principale(self):
        self.earn('50000', source=self.salaire)
        self.earn('150000', source=self.vente)
        top = services.top_source(self.user, services.period_bounds('month', self.today))
        self.assertEqual(top['name'], 'Vente')
        self.assertAlmostEqual(top['share'], 75.0)

    def test_solde_positif(self):
        self.earn('200000')
        self.spend('50000')
        solde = services.balance(self.user, services.period_bounds('month', self.today))
        self.assertEqual(solde['balance'], Decimal('150000'))
        self.assertTrue(solde['is_positive'])
        self.assertAlmostEqual(solde['spent_ratio'], 25.0)

    def test_solde_negatif(self):
        self.earn('30000')
        self.spend('45000')
        solde = services.balance(self.user, services.period_bounds('month', self.today))
        self.assertEqual(solde['balance'], Decimal('-15000'))
        self.assertFalse(solde['is_positive'])

    def test_solde_sans_revenu_ne_calcule_pas_de_ratio(self):
        self.spend('5000')
        solde = services.balance(self.user, services.period_bounds('month', self.today))
        self.assertIsNone(solde['spent_ratio'])
        self.assertEqual(solde['balance'], Decimal('-5000'))

    def test_le_reste_du_mois_passe_est_reporte(self):
        # Le reste anterieur (100000 - 60000 = 40000) est reporte sur le mois.
        last_month = date(2026, 6, 20)
        self.earn('100000', when=last_month)
        self.spend('60000', when=last_month)
        self.earn('50000')   # ce mois
        self.spend('20000')  # ce mois
        solde = services.balance(self.user, services.period_bounds('month', self.today))
        self.assertEqual(solde['carry_over'], Decimal('40000'))
        self.assertEqual(solde['income'], Decimal('50000'))
        self.assertEqual(solde['expense'], Decimal('20000'))
        self.assertEqual(solde['available'], Decimal('70000'))
        self.assertTrue(solde['is_positive'])

    def test_un_deficit_reporte_rend_le_solde_negatif(self):
        # Meme si le mois est positif, un report negatif peut rendre le solde negatif.
        self.spend('80000', when=date(2026, 6, 10))  # reserves entamees le mois passe
        self.earn('30000')
        self.spend('10000')
        solde = services.balance(self.user, services.period_bounds('month', self.today))
        self.assertEqual(solde['carry_over'], Decimal('-80000'))
        self.assertEqual(solde['available'], Decimal('-60000'))
        self.assertFalse(solde['is_positive'])

    def test_gagner_plus_est_une_amelioration(self):
        # A l'inverse des depenses : pour un revenu, monter est bon signe.
        self.earn('100000', when=self.today - timedelta(days=1))
        self.earn('150000', when=self.today)
        result = services.compare(self.user, 'day', self.today, model=Income)
        self.assertTrue(result['improving'])
        self.assertAlmostEqual(result['variation'], 50.0)

    def test_un_flux_recurrent_n_apparait_qu_une_fois(self):
        # Trois mois de salaire = un seul flux, pas trois lignes identiques.
        for mois in range(3):
            Income.objects.create(
                user=self.user, source=self.salaire, amount=Decimal('185000'),
                description='Salaire mensuel', is_recurring=True,
                date=self.today - timedelta(days=30 * mois),
            )
        flux = services.recurring_streams(self.user)
        self.assertEqual(len(flux), 1)
        self.assertEqual(flux[0].date, self.today)  # la plus recente

    def test_le_gabarit_du_solde_ne_laisse_pas_fuir_de_commentaire(self):
        self.earn('1000')
        response = self.client.get(reverse('income_dashboard'))
        self.assertNotContains(response, 'Partage entre l')

    def test_ajout_de_revenu(self):
        response = self.client.post(reverse('income_create'), {
            'amount': '185000',
            'source': self.salaire.pk,
            'date': self.today.isoformat(),
            'method': 'transfer',
            'description': 'Salaire de juillet',
            'note': '',
        })
        self.assertRedirects(response, reverse('income_dashboard'))
        self.assertEqual(Income.objects.count(), 1)

    def test_filtre_par_source(self):
        self.earn('10000', source=self.salaire)
        self.earn('20000', source=self.vente)
        response = self.client.get(reverse('income_history'), {'source': self.vente.pk})
        self.assertEqual(response.context['count'], 1)
        self.assertEqual(response.context['total'], Decimal('20000'))

    def test_la_source_utilisee_est_archivee_et_non_supprimee(self):
        self.earn('10000', source=self.salaire)
        self.client.post(reverse('source_delete', args=[self.salaire.pk]))
        self.salaire.refresh_from_db()
        self.assertTrue(self.salaire.is_archived)
        self.assertEqual(Income.objects.count(), 1)

    def test_les_revenus_sont_isoles_par_utilisateur(self):
        autre = User.objects.create_user('bob', password='motdepasse-123')
        Income.objects.create(
            user=autre, source=IncomeSource.objects.get(user=autre, name='Salaire'),
            amount=Decimal('777777'), date=self.today,
        )
        response = self.client.get(reverse('income_history'))
        self.assertNotContains(response, '777 777')
        self.assertEqual(response.context['count'], 0)

    def test_le_tableau_de_bord_affiche_le_solde(self):
        # Le tableau de bord utilise la date reelle : creer les donnees dessus.
        now = services.today()
        self.earn('100000', when=now)
        self.spend('40000', when=now)
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.context['balance']['balance'], Decimal('60000'))

    def test_les_pages_revenus_repondent(self):
        self.earn('50000')
        for name in ['income_dashboard', 'income_create', 'income_history',
                     'source_list']:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)

    def test_les_pages_revenus_exigent_une_connexion(self):
        self.client.logout()
        for name in ['income_dashboard', 'income_create', 'income_history',
                     'source_list']:
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 302, name)
            self.assertIn('/connexion/', response['Location'], name)

    def test_les_revenus_figurent_dans_l_export_excel(self):
        self.earn('123456')
        response = self.client.get(reverse('export_excel'))
        self.assertEqual(response.status_code, 200)
        from io import BytesIO

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(response.content))
        self.assertIn('Revenus', wb.sheetnames)
        valeurs = [c.value for row in wb['Revenus'].iter_rows() for c in row]
        self.assertIn(123456, valeurs)


class BudgetLimitViewTests(BaseCase):
    """La page Budgets ne doit jamais creer de limite invisible."""

    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)

    def _creer(self, **extra):
        donnees = {'period': 'week', 'amount': '20000', 'category': ''}
        donnees.update(extra)
        return self.client.post(reverse('budget_list'), donnees)

    def test_une_limite_creee_est_active_et_visible(self):
        self._creer(is_active='on')
        limite = BudgetLimit.objects.get(user=self.user, period='week')
        self.assertTrue(limite.is_active)
        cibles = [s['limit'] for s in services.all_limit_statuses(self.user)]
        self.assertIn(limite, cibles)

    def test_une_limite_desactivee_reste_visible_sur_la_page_budgets(self):
        # Sans cette regle, la limite serait invisible tout en bloquant la
        # creation d'une limite identique.
        self._creer()  # case « active » decochee
        limite = BudgetLimit.objects.get(user=self.user, period='week')
        self.assertFalse(limite.is_active)
        response = self.client.get(reverse('budget_list'))
        self.assertContains(response, 'desactivee')
        self.assertIn(limite, [s['limit'] for s in response.context['statuses']])

    def test_le_tableau_de_bord_ignore_les_limites_desactivees(self):
        self._creer()
        statuts = services.all_limit_statuses(self.user)
        self.assertEqual(statuts, [])

    def test_le_doublon_signale_une_limite_desactivee(self):
        self._creer()
        response = self._creer(amount='999')
        self.assertContains(response, 'desactivee existe deja')
        self.assertEqual(BudgetLimit.objects.filter(user=self.user).count(), 1)

    def test_la_case_active_est_precochee_dans_le_formulaire(self):
        response = self.client.get(reverse('budget_list'))
        self.assertContains(response, 'name="is_active"')
        self.assertContains(response, 'checked')


class OverspendAlertTests(BaseCase):
    """Alerte quand les depenses du mois depassent les revenus du mois."""

    def setUp(self):
        super().setUp()
        self.salaire = IncomeSource.objects.get(user=self.user, name='Salaire')

    def earn(self, amount, when=None):
        return Income.objects.create(
            user=self.user, source=self.salaire,
            amount=Decimal(amount), date=when or self.today,
        )

    def _overspend(self):
        return Alert.objects.filter(user=self.user, kind=Alert.Kind.OVERSPEND)

    def test_pas_d_alerte_quand_les_revenus_couvrent_les_depenses(self):
        self.earn('100000')
        self.spend('60000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().count(), 0)

    def test_alerte_quand_les_depenses_depassent_les_revenus(self):
        self.earn('100000')
        self.spend('130000')
        services.evaluate_alerts(self.user, self.today)
        alert = self._overspend().get()
        self.assertEqual(alert.level, Alert.Level.EXCEEDED)
        self.assertEqual(alert.spent, Decimal('130000'))
        self.assertEqual(alert.limit_amount, Decimal('100000'))
        self.assertIsNone(alert.limit_id)

    def test_pas_d_alerte_sans_revenu_enregistre(self):
        # Un utilisateur qui ne suit que ses depenses ne doit pas etre alarme.
        self.spend('50000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().count(), 0)

    def test_l_alerte_n_est_pas_dupliquee(self):
        self.earn('100000')
        self.spend('130000')
        services.evaluate_alerts(self.user, self.today)
        self.spend('5000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().count(), 1)

    def test_l_alerte_non_lue_se_met_a_jour(self):
        self.earn('100000')
        self.spend('130000')
        services.evaluate_alerts(self.user, self.today)
        self.spend('20000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().get().spent, Decimal('150000'))

    def test_l_alerte_non_lue_disparait_si_la_situation_se_retablit(self):
        self.earn('100000')
        self.spend('130000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().count(), 1)
        self.earn('60000')  # les revenus repassent au-dessus des depenses
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().count(), 0)

    def test_une_alerte_lue_reste_comme_trace(self):
        self.earn('100000')
        self.spend('130000')
        services.evaluate_alerts(self.user, self.today)
        self._overspend().update(is_read=True)
        self.earn('60000')
        services.evaluate_alerts(self.user, self.today)
        # Deja acquittee : on la conserve comme historique plutot que l'effacer.
        self.assertEqual(self._overspend().count(), 1)

    def test_ajouter_un_revenu_resout_l_alerte_via_la_vue(self):
        self.client.force_login(self.user)
        self.spend('130000')
        self.earn('100000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(self._overspend().count(), 1)
        self.client.post(reverse('income_create'), {
            'amount': '60000', 'source': self.salaire.pk,
            'date': self.today.isoformat(), 'method': 'transfer',
            'description': '', 'note': '',
        })
        self.assertEqual(self._overspend().count(), 0)

    def test_l_alerte_apparait_sur_le_tableau_de_bord(self):
        # Le tableau de bord utilise la date reelle : creer les donnees dessus.
        now = services.today()
        self.client.force_login(self.user)
        self.earn('100000', when=now)
        self.spend('130000', when=now)
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, 'Depenses superieures aux revenus')

    def test_l_alerte_compte_dans_le_badge_non_lu(self):
        self.earn('100000')
        self.spend('130000')
        services.evaluate_alerts(self.user, self.today)
        self.assertEqual(len(services.unread_alerts(self.user)), 1)


class ReportTests(BaseCase):
    def test_le_rapport_est_genere_selon_la_date_des_depenses(self):
        self.spend('2000', category=self.food)
        self.spend('3000', category=self.transport)
        report = services.build_report(self.user, 'month', self.today)
        self.assertEqual(report.total, Decimal('5000'))
        self.assertEqual(report.expense_count, 2)
        self.assertEqual(report.top_category, 'Transport')

    def test_le_rapport_est_rafraichi_et_non_duplique(self):
        self.spend('2000')
        services.build_report(self.user, 'month', self.today)
        self.spend('1000')
        report = services.build_report(self.user, 'month', self.today)
        self.assertEqual(Report.objects.filter(kind='month').count(), 1)
        self.assertEqual(report.total, Decimal('3000'))


class ViewTests(BaseCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)

    def test_le_tableau_de_bord_repond(self):
        self.spend('1200')
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Depenses d')

    def test_ajout_de_depense(self):
        response = self.client.post(reverse('expense_create'), {
            'amount': '2500',
            'category': self.food.pk,
            'date': self.today.isoformat(),
            'payment_method': 'cash',
            'description': 'Dejeuner',
            'note': '',
        })
        self.assertRedirects(response, reverse('dashboard'))
        self.assertEqual(Expense.objects.count(), 1)

    def test_le_champ_date_est_prerempli_au_format_iso(self):
        # <input type="date"> reste vide si la valeur n'est pas en ISO.
        response = self.client.get(reverse('expense_create'))
        attendu = f'value="{services.today().isoformat()}"'
        self.assertContains(response, attendu)

    def test_les_donnees_sont_isolees_par_utilisateur(self):
        autre = User.objects.create_user('bob', password='motdepasse-123')
        Expense.objects.create(
            user=autre, category=Category.objects.get(user=autre, name='Sante'),
            amount=Decimal('9999'), date=self.today,
        )
        response = self.client.get(reverse('history'))
        self.assertNotContains(response, '9 999')

    def test_filtre_par_categorie(self):
        self.spend('1000', category=self.food)
        self.spend('2000', category=self.transport)
        response = self.client.get(reverse('history'), {'category': self.transport.pk})
        self.assertEqual(response.context['count'], 1)
        self.assertEqual(response.context['total'], Decimal('2000'))

    def test_filtre_par_montant(self):
        self.spend('100')
        self.spend('5000')
        response = self.client.get(reverse('history'), {'min_amount': '1000'})
        self.assertEqual(response.context['count'], 1)

    def test_la_categorie_utilisee_est_archivee_et_non_supprimee(self):
        self.spend('500', category=self.food)
        self.client.post(reverse('category_delete', args=[self.food.pk]))
        self.food.refresh_from_db()
        self.assertTrue(self.food.is_archived)
        self.assertEqual(Expense.objects.count(), 1)

    def test_la_categorie_inutilisee_est_supprimee(self):
        self.client.post(reverse('category_delete', args=[self.transport.pk]))
        self.assertFalse(Category.objects.filter(pk=self.transport.pk).exists())

    def test_export_excel(self):
        self.spend('1000')
        response = self.client.get(reverse('export_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

    def test_export_pdf(self):
        self.spend('1000')
        response = self.client.get(reverse('export_pdf'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_les_pages_exigent_une_connexion(self):
        self.client.logout()
        for name in ['dashboard', 'history', 'category_list', 'budget_list',
                     'reports', 'settings', 'expense_create']:
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 302, name)
            self.assertIn('/connexion/', response['Location'], name)

    def test_les_pages_principales_repondent(self):
        BudgetLimit.objects.create(user=self.user, period='week', amount=Decimal('5000'))
        self.spend('4500')
        for name in ['dashboard', 'history', 'category_list', 'budget_list',
                     'reports', 'settings', 'expense_create']:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)

    def test_inscription(self):
        self.client.logout()
        response = self.client.post(reverse('signup'), {
            'username': 'charlie',
            'email': '',
            'password1': 'un-mot-de-passe-solide-42',
            'password2': 'un-mot-de-passe-solide-42',
        })
        self.assertRedirects(response, reverse('dashboard'))
        charlie = User.objects.get(username='charlie')
        self.assertEqual(Category.objects.filter(user=charlie).count(), 11)
