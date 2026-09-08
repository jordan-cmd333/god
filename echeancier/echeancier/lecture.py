"""Lecture du classeur Excel : reperage des colonnes et extraction des taches."""

from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .modele import Tache, normaliser

# Synonymes acceptes pour chaque role de colonne. La comparaison se fait sur
# le libelle normalise (minuscules, sans accent) et par prefixe.
SYNONYMES: dict[str, tuple[str, ...]] = {
    "echeance": (
        "echeance", "date echeance", "date d echeance", "date limite", "limite",
        "date limite d execution", "date de fin", "date fin", "fin prevue",
        "a rendre le", "deadline", "due date", "due", "date butoir", "butoir",
        "date cible", "livraison", "date de livraison",
    ),
    "tache": (
        "tache", "taches", "libelle", "intitule", "designation", "description",
        "objet", "action", "activite", "sujet", "titre", "mission", "item",
        "task", "todo", "a faire",
    ),
    "statut": (
        "statut", "status", "etat", "avancement", "progression", "etape",
        "realise", "fait", "termine", "cloture",
    ),
    "responsable": (
        "responsable", "pilote", "assigne", "assignee", "affecte a", "charge",
        "qui", "acteur", "owner", "porteur", "referent", "intervenant",
    ),
    "priorite": ("priorite", "urgence", "criticite", "importance", "priority"),
    "identifiant": ("id", "ref", "reference", "n", "no", "num", "numero", "code"),
}

FORMATS_DATE = (
    "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y",
    "%m/%d/%Y", "%Y/%m/%d", "%d %m %Y",
)

MOIS_FR = {
    "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6,
    "juillet": 7, "aout": 8, "septembre": 9, "octobre": 10, "novembre": 11,
    "decembre": 12,
}

ORIGINE_EXCEL = datetime(1899, 12, 30)
LIGNES_SCRUTEES = 15


class LectureError(Exception):
    pass


def _role_de(entete: object) -> str | None:
    """Role de colonne devine a partir d'un libelle d'en-tete."""
    texte = normaliser(entete)
    if not texte:
        return None
    for role, mots in SYNONYMES.items():
        if texte in mots:
            return role
    for role, mots in SYNONYMES.items():
        for mot in mots:
            if len(mot) >= 4 and (texte.startswith(mot) or mot in texte):
                return role
    return None


def convertir_date(valeur: object) -> date | None:
    """Accepte les cellules date, les series Excel et les dates saisies en texte."""
    if valeur is None:
        return None
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    if isinstance(valeur, bool):
        return None
    if isinstance(valeur, (int, float)):
        if not 1 <= float(valeur) <= 200000:
            return None
        return (ORIGINE_EXCEL + timedelta(days=float(valeur))).date()

    texte = str(valeur).strip()
    if not texte:
        return None
    for gabarit in FORMATS_DATE:
        try:
            return datetime.strptime(texte, gabarit).date()
        except ValueError:
            continue
    # « 12 mars 2026 » et variantes.
    morceaux = normaliser(texte).split()
    if len(morceaux) == 3 and morceaux[1] in MOIS_FR:
        try:
            return date(int(morceaux[2]), MOIS_FR[morceaux[1]], int(morceaux[0]))
        except ValueError:
            return None
    return None


def _texte(valeur: object) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return valeur.date().isoformat()
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def reperer_entete(
    lignes: list[tuple], forcees: dict[str, str] | None = None
) -> tuple[int, dict[str, int]] | None:
    """Repere la ligne d'en-tete et l'index de chaque colonne utile.

    Renvoie (index de la ligne d'en-tete, {role: index de colonne}) ou None.
    """
    forcees = {role: normaliser(nom) for role, nom in (forcees or {}).items()}
    meilleur: tuple[int, int, dict[str, int]] | None = None

    for index, ligne in enumerate(lignes[:LIGNES_SCRUTEES]):
        roles: dict[str, int] = {}
        for colonne, cellule in enumerate(ligne):
            libelle = normaliser(cellule)
            if not libelle:
                continue
            role = next((r for r, nom in forcees.items() if nom == libelle), None)
            if role is None:
                if libelle in forcees.values():
                    continue
                role = _role_de(cellule)
            if role and role not in roles:
                roles[role] = colonne
        if "echeance" not in roles:
            continue
        if "tache" not in roles:
            secours = _colonne_texte(lignes, index, exclure=set(roles.values()))
            if secours is None:
                continue
            roles["tache"] = secours
        score = len(roles)
        if meilleur is None or score > meilleur[1]:
            meilleur = (index, score, roles)

    if meilleur is None:
        return None
    return meilleur[0], meilleur[2]


def _colonne_texte(lignes: list[tuple], entete: int, exclure: set[int]) -> int | None:
    """A defaut de colonne « tache », prend la premiere colonne de texte peuplee."""
    largeur = max((len(l) for l in lignes[entete + 1 : entete + 12]), default=0)
    for colonne in range(largeur):
        if colonne in exclure:
            continue
        valeurs = [
            _texte(ligne[colonne])
            for ligne in lignes[entete + 1 : entete + 12]
            if colonne < len(ligne)
        ]
        peuplees = [v for v in valeurs if v]
        if len(peuplees) >= max(1, len(valeurs) // 2) and any(
            convertir_date(v) is None for v in peuplees
        ):
            return colonne
    return None


def lire_classeur(
    chemin: Path,
    feuilles: list[str] | None = None,
    colonnes_forcees: dict[str, str] | None = None,
) -> tuple[list[Tache], dict[str, dict[str, str]]]:
    """Renvoie les taches du classeur et, par feuille, les colonnes reconnues."""
    if not chemin.is_file():
        raise LectureError(f"Classeur introuvable : {chemin}")
    if chemin.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise LectureError(
            f"Format non pris en charge ({chemin.suffix}). Enregistrez le classeur "
            "au format .xlsx depuis Excel."
        )

    try:
        classeur = load_workbook(chemin, read_only=True, data_only=True)
    except Exception as erreur:  # fichier corrompu, verrouille, protege...
        raise LectureError(f"Lecture impossible du classeur : {erreur}") from erreur

    toutes: list[Tache] = []
    diagnostic: dict[str, dict[str, str]] = {}
    try:
        for feuille in classeur.worksheets:
            if feuilles and feuille.title not in feuilles:
                continue
            lignes = [tuple(l) for l in feuille.iter_rows(values_only=True)]
            entete = reperer_entete(lignes, colonnes_forcees)
            if entete is None:
                diagnostic[feuille.title] = {}
                continue
            index_entete, roles = entete
            diagnostic[feuille.title] = {
                role: _texte(lignes[index_entete][col]) or f"colonne {col + 1}"
                for role, col in sorted(roles.items(), key=lambda kv: kv[1])
            }
            toutes.extend(_taches_de(feuille.title, lignes, index_entete, roles))
    finally:
        classeur.close()

    return toutes, diagnostic


def _taches_de(
    nom: str, lignes: list[tuple], index_entete: int, roles: dict[str, int]
) -> list[Tache]:
    taches: list[Tache] = []
    for numero, ligne in enumerate(lignes[index_entete + 1 :], start=index_entete + 2):
        def cellule(role: str) -> object:
            colonne = roles.get(role)
            if colonne is None or colonne >= len(ligne):
                return None
            return ligne[colonne]

        libelle = _texte(cellule("tache"))
        echeance = convertir_date(cellule("echeance"))
        if not libelle and echeance is None:
            continue
        taches.append(
            Tache(
                libelle=libelle or f"(sans libelle, ligne {numero})",
                echeance=echeance,
                feuille=nom,
                ligne=numero,
                identifiant=_texte(cellule("identifiant")),
                statut=_texte(cellule("statut")),
                responsable=_texte(cellule("responsable")),
                priorite=_texte(cellule("priorite")),
            )
        )
    return taches
