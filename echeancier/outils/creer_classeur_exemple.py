"""Genere un classeur Excel de demonstration pour essayer l'echeancier.

Usage : python3 outils/creer_classeur_exemple.py [chemin.xlsx]
"""

from __future__ import annotations

import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

LIGNES = [
    ("T-001", "Rendre le rapport trimestriel", -3, "En cours", "Sophie", "Haute"),
    ("T-002", "Relancer le fournisseur Duval", -1, "A faire", "Karim", "Haute"),
    ("T-003", "Valider le budget formation", 0, "En cours", "Sophie", "Haute"),
    ("T-004", "Signer l'avenant du bail", 1, "A faire", "Marc", "Haute"),
    ("T-005", "Preparer la reunion clients", 3, "A faire", "Karim", "Moyenne"),
    ("T-006", "Mettre a jour l'inventaire", 6, "A faire", "Lea", "Basse"),
    ("T-007", "Commander les fournitures", 12, "A faire", "Lea", "Basse"),
    ("T-008", "Archiver les dossiers 2025", 30, "A faire", "Marc", "Basse"),
    ("T-009", "Declaration sociale annuelle", -10, "Termine", "Sophie", "Haute"),
    ("T-010", "Entretien annuel de l'equipe", 2, "Fait", "Marc", "Moyenne"),
]

ENTETES = ["Ref", "Tache", "Date limite", "Statut", "Responsable", "Priorite"]


def creer(chemin: Path) -> None:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Suivi des taches"

    feuille["A1"] = "Suivi des taches du service"   # ligne parasite volontaire :
    feuille["A2"] = ""                              # l'en-tete n'est pas en ligne 1
    feuille.append([])
    for colonne, entete in enumerate(ENTETES, start=1):
        cellule = feuille.cell(row=3, column=colonne, value=entete)
        cellule.font = Font(bold=True)

    aujourdhui = date.today()
    for ref, libelle, decalage, statut, qui, priorite in LIGNES:
        feuille.append(
            [ref, libelle, aujourdhui + timedelta(days=decalage), statut, qui, priorite]
        )

    for colonne in range(1, len(ENTETES) + 1):
        feuille.column_dimensions[get_column_letter(colonne)].width = 26
    for ligne in feuille.iter_rows(min_row=4, min_col=3, max_col=3):
        for cellule in ligne:
            cellule.number_format = "DD/MM/YYYY"

    classeur.save(chemin)


if __name__ == "__main__":
    destination = Path(sys.argv[1] if len(sys.argv) > 1 else "exemple-taches.xlsx")
    creer(destination)
    print(f"Classeur d'exemple cree : {destination.resolve()}")
